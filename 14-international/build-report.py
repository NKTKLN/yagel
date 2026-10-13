from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'Ягель — отчёт по практической №14.docx'
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(1.8), Cm(1.7)
s.left_margin, s.right_margin = Cm(2.2), Cm(1.8)
s.header_distance, s.footer_distance = Cm(.7), Cm(.7)
s.different_first_page_header_footer = True
normal = d.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 12)]:
    st = d.styles[name]
    st.font.name = 'Times New Roman'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(9)
    st.paragraph_format.space_after = Pt(6)

footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def p(text, bold=False, size=None, center=False):
    para = d.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    r = para.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    return para


def h(text):
    d.add_heading(text, level=1)


def sub(text):
    d.add_heading(text, level=2)


def source(num, title, url=None, detail=''):
    para = d.add_paragraph()
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.0
    r = para.add_run(f'[{num}] ')
    r.font.size = Pt(10.5)
    if url:
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '000000')
        props.append(color)
        underline = OxmlElement('w:u')
        underline.set(qn('w:val'), 'single')
        props.append(underline)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '21')
        props.append(sz)
        run.append(props)
        text = OxmlElement('w:t')
        text.text = title
        run.append(text)
        link.append(run)
        para._p.append(link)
    else:
        para.add_run(title).font.size = Pt(10.5)
    if detail:
        para.add_run(' ' + detail).font.size = Pt(10.5)


# 1. Титульный лист.
p('МИНОБРНАУКИ РОССИИ', center=True, size=12)
p('МИРЭА — Российский технологический университет', center=True, bold=True)
p('Институт перспективных технологий\nи индустриального программирования', center=True)
p('Кафедра индустриального программирования', center=True)
for _ in range(3):
    p('')
p('ОТЧЁТ', bold=True, size=22, center=True)
p('по практическому занятию № 14', bold=True, size=16, center=True)
p('по дисциплине «Стартап в информационных технологиях»', center=True)
p('Планирование и осуществление международного расширения стартапа, изучение международных рынков', size=15, center=True)
p('Проект «Ягель A1»\nВыход на рынок Германии', bold=True, size=15, center=True)
for _ in range(2):
    p('')
q = p('Выполнил: Калинин Никита Викторович\nГруппа: ЭФБО-09-23\nПреподаватель: Адышкин Сергей Сергеевич')
q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p('')
p('Москва, 2026', center=True)
d.add_page_break()

# 2. Проект и рынок.
h('1. Проект и выбор рынка')
p('Цель работы — предложить план выхода «Ягеля A1» на рынок Германии. Основная аудитория — владельцы локального умного дома. В отчёте рассмотрены выбор рынка, адаптация продукта и порядок запуска.')
sub('Что представляет собой продукт')
p('«Ягель A1» — настольные часы на электронных чернилах и монитор воздуха. Он показывает CO₂, температуру, влажность и давление. Прототип использует ESP32-C3, датчики MH-Z19B и BME280, часы DS3231 и экран 2,9″. Схема, прошивка и модель корпуса открыты; в проекте указана лицензия MIT [1].')
p('Главная идея — работа без обязательного облака и серверов производителя. Прибор измеряет параметры и показывает их на собственном экране. В целевой версии он также передаёт данные в домашнюю сеть, без регистрации в сервисе производителя и подписки.', bold=True)
p('Владелец сможет использовать Home Assistant на своём оборудовании для истории и автоматизаций. Интернет нужен только для добровольных функций, например синхронизации часов. При его отключении основные возможности должны сохраняться.')
sub('Что уже сделано')
p('В проверенной прошивке реализованы измерения, экран и чтение времени из DS3231. Интернет-синхронизация часов сейчас включена по умолчанию. Локальный HTTP API, MQTT и автоматическое подключение к Home Assistant пока предстоит реализовать. В отчёте они рассматриваются как план развития [1].')
sub('Почему выбрана Германия — задание № 1')
p('По опросу Bitkom, в 2026 году 52% жителей Германии от 16 лет используют дома устройства Smart Home; в 2024 году — 46%. Для 87% пользователей важны безопасность данных и приватность, для 79% — совместимость [2]. Поэтому Германия подходит для проверки устройства, которое подключается к существующему умному дому и оставляет данные у владельца.')
p('Конкуренция уже есть: Aranet4 работает автономно и использует Bluetooth, а открытый AirGradient ONE поддерживает локальный Home Assistant [3, 4]. «Ягель» предлагает сочетание часов на E-Ink, открытой конструкции и локального доступа. Объём продаж именно локальных CO₂-мониторов в Германии пока не подтверждён; его нельзя приравнивать ко всему рынку Smart Home.')
p('NBS рекомендует готовить для переговоров аргументы об общественной пользе продукта [17]. Для «Ягеля» это помощь в проветривании, возможность разумнее расходовать энергию и продлевать срок службы электроники. Эти преимущества нужно показать на конкретных примерах и подтвердить испытаниями.')
d.add_page_break()

# 3. Польза и адаптация.
h('2. Польза и адаптация продукта — задание № 2')
sub('Здоровье и условия в помещении')
p('Монитор помогает заметить рост CO₂ и вовремя проветрить комнату. Umweltbundesamt рассматривает CO₂ как показатель достаточности вентиляции; превышение 1 000 ppm — повод улучшить воздухообмен [14]. «Ягель» помогает поддерживать подходящие условия дома, но сам не очищает воздух и не является медицинским прибором.')
sub('Энергоэффективность')
p('Показания помогают проветривать по необходимости, а через Home Assistant — управлять совместимой вентиляцией. Это может уменьшить лишние теплопотери. Umweltbundesamt рекомендует правильное проветривание как часть экономного отопления [15]. Экономию именно от «Ягеля» нужно проверить в пилоте. Потребление всего прибора также нужно измерить: оно зависит не только от E-Ink, но и от датчиков и Wi-Fi.')
sub('Экология и срок службы')
p('Открытые схема, код и корпус создают основу для ремонта и замены деталей. Отсутствие обязательного облака позволяет сохранить основные функции при прекращении работы сервиса производителя. Для практической пользы нужны доступные запчасти и инструкция по ремонту. Продление срока использования электроники помогает беречь ресурсы [16]; экологический след самого «Ягеля» ещё не рассчитан.')
sub('Что изменить перед запуском')
p('Интерфейс и документация. Подготовить немецкую инструкцию и страницу настройки, формат времени 24 часа, часовой пояс Германии с переходом на летнее время. Сохранить привычные единицы: °C, %, hPa и ppm.')
p('Локальная работа. Реализовать HTTP API и MQTT с обнаружением в Home Assistant. Добавить настройку Wi-Fi без перепрошивки, установку времени локально и добровольную синхронизацию через интернет. Защитить настройки и предусмотреть установку проверенных обновлений из локальной сети.')
p('Продвижение и поддержка. Показать три сценария: своевременное проветривание, управление вентиляцией и ремонт. Немецкий партнёр отвечает за первую линию поддержки и местные возвраты; команда — за прошивку и сложные вопросы. Инструкции должны быть доступны для скачивания.')
sub('Цена')
p('Проверяемая гипотеза — 129 € с НДС, без подписки. При ставке НДС 19% продавцу остаётся 108,40 € до расходов [5]. Из этой суммы оплачиваются производство, доля партнёра, доставка, возвраты и поддержка. До запуска нужно получить полную смету, отдельно учесть испытания и регистрации и проверить готовность покупателей платить эту цену.')
d.add_page_break()

# 4. План и риски.
h('3. План выхода и риски — задание № 3')
p('Выход планируется через немецкого продавца или интегратора. Команда разрабатывает прибор и отвечает за качество. Партнёр продаёт его, организует местную доставку и обслуживание; обязанности импортёра закрепляются договором. Сроки ниже — ориентиры. К следующему этапу переходят после выполнения условий предыдущего.')
sub('Этапы запуска')
p('1. Проверить законность поставок — месяц 1. Определить таможенный код и происхождение товара, проверить санкции, банки, участников сделки и перевозку. Получить заключение о допустимости поставок и расчётов. Партнёр в Германии сам по себе не снимает ограничений [6].')
p('2. Проверить спрос — месяцы 1–2. Провести 15 интервью, проверить цену 129 € и найти партнёра. Ориентир — минимум 10 заявок на пилот без оплаты. До переговоров подготовить немецкое описание общественной пользы: проветривание, план проверки энергосбережения, ремонт и независимость от облака [17].')
p('3. Подготовить продукт — месяцы 2–4. Завершить локальные функции и документацию, исправить крепления корпуса, проверить CO₂ по эталонному прибору. Измерить потребление, описать замену деталей. Выполнить применимые испытания и регистрации до передачи устройств пользователям.')
p('4. Провести пилот — месяцы 4–5. Передать 30 приборов после правового допуска. Цели: минимум 27 самостоятельных подключений и 72 часа локальной работы без интернета. Собрать отзывы. Экономию оценивать по расходу энергии при сходных условиях и качестве воздуха.')
p('5. Решить вопрос о росте — с месяца 6. Исправить проблемы, проверить возвраты и фактическую прибыль с устройства. Расширять продажи после подтверждения спроса и качества. Продвижение — через сообщества умного дома, обзоры и демонстрации партнёра.')
sub('Правовые условия и риски')
p('До продаж нужно подтвердить соответствие требованиям RED и RoHS, подготовить документы и маркировку CE [7, 8]. Определить обязанности по электроотходам и упаковке: WEEE/ElektroG и LUCID. Иностранному производителю без присутствия в Германии для WEEE-регистрации нужен местный уполномоченный [9, 10].')
p('Нужно учесть кибербезопасность: применимые требования RED и CRA; отчётность CRA действует с 11.09.2026, основные обязанности — с 11.12.2027 [7, 11]. У продавца должны быть предусмотрены 14-дневный отказ при дистанционной продаже и минимум 2 года законной ответственности за недостатки; данные заказов и обращений обрабатываются по GDPR [12, 13].')
p('Главные риски — запрет поставки, высокие затраты, ошибки измерений и слабый спрос. Их снижают проверкой сделки, полной сметой, испытаниями и небольшой партией. Пользу продукта подтверждают результатами пилота. При запрете поставок запуск останавливают.')
d.add_page_break()

# 5. Выводы и проверяемые источники.
h('4. Выводы')
p('«Ягель A1» выходит на немецкий рынок как открытый прибор для более осознанного проветривания и контроля условий дома. Общественная польза связана с заботой о качестве воздуха, возможностью разумнее управлять вентиляцией и продлевать срок службы устройства. Работа без обязательного облака сохраняется как основа продукта.')
p('Первый шаг — проверить спрос и законность поставок. Затем подготовить локальные функции, немецкую документацию и пилот из 30 приборов через местного партнёра. Цена 129 € и экономия энергии остаются гипотезами до проверки. Увеличивать продажи следует после подтверждения качества, спроса и экономики.')
h('5. Источники')
p('Веб-источники проверены 13.10.2026. Ссылки открываются нажатием на название. Числа интервью, заявок, сроки и критерии пилота — плановые ориентиры, а не результаты проведённых исследований.', size=10.5)
source(1, 'Материалы проекта и исходники eink-environment-clock.', 'https://github.com/NKTKLN/eink-environment-clock', 'Локально: README.md, firmware/src/app_main.c, app_config.h, net_time.c, LICENSE; паспорт проекта в accelerator/.')
source(2, 'Bitkom. Smart Home: Vom WLAN-Thermostat zum menschenähnlichen Roboter.', 'https://www.bitkom.org/Presse/Presseinformation/Smart-Home-vom-Thermostat-zum-Roboter', '24.08.2026. Опрос, методика и показатели рынка.')
source(3, 'Aranet. Aranet4 HOME — FAQ.', 'https://help.aranet.com/aranet4/aranet4-home/faq', 'Автономная работа и Bluetooth.')
source(4, 'AirGradient. Indoor Air Quality Monitor / ONE.', 'https://www.airgradient.com/indoor/', 'Открытость и локальная интеграция Home Assistant.')
source(5, 'BMF. Umsatzsteuer, § 12 — Steuersätze.', 'https://usth.bundesfinanzministerium.de/usth/2024/A-Umsatzsteuergesetz/IV-Steuer-und-Vorsteuer/Paragraf-12/inhalt.html', 'Стандартная ставка 19%.')
source(6, 'Совет ЕС. EU sanctions against Russia: questions and answers.', 'https://www.consilium.europa.eu/en/policies/sanctions-against-russia-explained/', 'Ограничения торговли, расчётов и противодействие обходу.')
source(7, 'Еврокомиссия. Radio Equipment Directive (RED).', 'https://single-market-economy.ec.europa.eu/sectors/electrical-and-electronic-engineering-industries-eei/radio-equipment-directive-red_en', 'Безопасность, радиочастоты, CE и кибербезопасность.')
source(8, 'Еврокомиссия. RoHS Directive.', 'https://environment.ec.europa.eu/topics/waste-and-recycling/rohs-directive_en', 'Ограничения опасных веществ в электронике.')
source(9, 'Stiftung ear. WEEE-Registrierung als ausländisches Unternehmen.', 'https://www.stiftung-ear.de/anleitungen/weee-nummer-beantragen-als-auslaendisches-unternehmen/', 'Регистрация и уполномоченный представитель.')
source(10, 'ZSVR. The Verpackungsgesetz and online retailers.', 'https://www.verpackungsregister.org/en/knowledge-bases/mail-order-companies-and-online-retailers', 'Упаковка, LUCID и обязанности участников поставки.')
source(11, 'Еврокомиссия. Cyber Resilience Act.', 'https://digital-strategy.ec.europa.eu/en/policies/cyber-resilience-act', 'Обязанности производителя и сроки применения.')
source(12, 'Your Europe. Consumer guarantees, warranties, claims and returns.', 'https://europa.eu/youreurope/business/selling-in-eu/consumer-contracts-guarantees/consumer-guarantees/index_en.htm', 'Правила дистанционной продажи и ответственность за недостатки.')
source(13, 'Еврокомиссия. Application of the GDPR.', 'https://commission.europa.eu/law/law-topic/data-protection/information-business-and-organisations/application-gdpr_en', 'Применимость к обработке персональных данных.')
source(14, 'Umweltbundesamt. Richtig Lüften in Schulen.', 'https://www.umweltbundesamt.de/richtig-lueften-in-schulen', 'CO₂ как показатель вентиляции, рекомендации по проветриванию.')
source(15, 'Umweltbundesamt. Richtiges Heizen schützt das Klima und den Geldbeutel.', 'https://www.umweltbundesamt.de/umwelttipps-fuer-den-alltag/richtiges-heizen-schuetzt-das-klima-den-geldbeutel', 'Отопление, проветривание и энергосбережение.')
source(16, 'Umweltbundesamt. Verlängerung der Produktnutzungsdauer.', 'https://www.umweltbundesamt.de/publikationen/verlaengerung-der-produktnutzungsdauer', '2020. Срок службы электроники и сохранение ресурсов.')
source(17, 'NBS. Немецкий рынок: где вход?', 'https://nbs-research.com/business/marketing/немецкий-рынок-где-вход', '15.06.2019. Практический пример аргументов для переговоров; не исследование всех покупателей.')

d.core_properties.title = 'Практическое занятие №14 — Международное расширение «Ягель A1»'
d.core_properties.subject = 'Германия: экология, энергоэффективность, здоровье и работа без обязательного облака'
d.core_properties.author = 'Калинин Никита Викторович'
d.core_properties.keywords = 'Ягель, YAGEL, Германия, локальный умный дом, практическая 14'
d.save(OUT)
print(OUT)
